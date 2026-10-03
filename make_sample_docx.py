from docx import Document

file_path = "sample.docx"

doc = Document()
doc.add_heading("Sample DOCX", 0)
doc.add_paragraph("Hello! This is a sample Word document for testing LangChain loader.")
doc.save(file_path)

print(f"DOCX created: {file_path}")
